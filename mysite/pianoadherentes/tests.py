from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import Adherente, Log, Titular


class BajasMasivasTests(TestCase):
	def setUp(self):
		self.usuario = User.objects.create_user(username='operador', password='clave')
		self.usuario.user_permissions.add(Permission.objects.get(codename='can_view_stats_controllerAdmin'))
		self.usuario.user_permissions.add(Permission.objects.get(codename='can_view_stats', content_type__model='titular'))
		self.creador = User.objects.create_user(username='creador', password='clave')
		self.client.force_login(self.usuario)
		self.titular = Titular.objects.create(
			cbu='0012345678901234567890', name='Ana', last_name='Perez',
			document_type='DNI', document='12345678', sex='F',
			street_address='Calle', number='1', province='Buenos Aires',
			city='Ciudad', postal_code='1000', user_upload=self.creador,
		)
		self.activo = self.adherente('Activo')
		self.otro_activo = self.adherente('Segundo')
		self.inactivo = self.adherente('Inactivo', is_active=False)
		self.otro_inactivo = self.adherente('Otro inactivo', is_active=False)

	def adherente(self, nombre, is_active=True):
		return Adherente.objects.create(
			titular=self.titular, name=nombre, last_name='Perez',
			document_type='DNI', document='12345678', sex='F',
			street_address='Calle', number='1', province='Buenos Aires',
			city='Ciudad', postal_code='1000', user_upload=self.creador,
			sucursal='Centro', legajo=1, is_active=is_active,
		)

	def subir(self, filas, vista_previa=False):
		libro = Workbook()
		for fila in filas:
			libro.active.append(fila)
		contenido = BytesIO()
		libro.save(contenido)
		archivo = SimpleUploadedFile('cbus.xlsx', contenido.getvalue())
		return self.client.post(reverse('bajas_masivas'), {
			'archivo': archivo, 'vista_previa': '1' if vista_previa else '',
		})

	def test_vista_previa_cuenta_registros_sin_dar_de_baja(self):
		respuesta = self.subir([[self.titular.cbu], [], ['9' * 22], [self.titular.cbu], ['cbu']], vista_previa=True)
		self.assertEqual(respuesta.status_code, 200)
		self.assertEqual(respuesta.json(), {'cantidad': 3, 'filas': 4})
		self.titular.refresh_from_db()
		self.activo.refresh_from_db()
		self.assertTrue(self.titular.is_active)
		self.assertTrue(self.activo.is_active)
		self.assertEqual(Log.objects.count(), 0)

	def test_vista_previa_vacia_no_procesa(self):
		respuesta = self.subir([[]], vista_previa=True)
		self.assertEqual(respuesta.status_code, 400)
		self.titular.refresh_from_db()
		self.assertTrue(self.titular.is_active)

	def test_procesa_y_reporta_todas_las_filas(self):
		respuesta = self.subir([
			[self.titular.cbu], ['9' * 22], [self.titular.cbu],
			['cbu'], [self.titular.cbu, 'dato extra'], [1234567890123456789012], ['=1+1'],
		])
		self.assertEqual(respuesta.status_code, 200)
		filas = list(load_workbook(BytesIO(respuesta.content)).active.values)
		self.assertEqual([fila[2] for fila in filas[1:]], [
			'Baja realizada', 'No encontrado', 'Ya estaban dados de baja',
			'CBU invalido', 'CBU invalido', 'CBU invalido', 'CBU invalido',
		])
		self.assertEqual(filas[1][1], self.titular.cbu)
		self.assertEqual(filas[1][4], 'Activo Perez, Segundo Perez')
		self.assertEqual(filas[1][5], 'Inactivo Perez, Otro inactivo Perez')
		self.assertEqual(filas[1][7], self.usuario.username)
		self.titular.refresh_from_db()
		self.activo.refresh_from_db()
		self.inactivo.refresh_from_db()
		self.assertFalse(self.titular.is_active)
		self.assertEqual(self.titular.user_upload, self.usuario)
		self.assertIsNotNone(self.titular.deleted)
		self.assertFalse(self.activo.is_active)
		self.assertEqual(self.activo.user_upload, self.usuario)
		self.assertIsNotNone(self.activo.deleted)
		self.assertEqual(self.inactivo.user_upload, self.creador)
		self.assertEqual(Log.objects.filter(adherente=self.activo, movimiento='Baja', user=self.usuario).count(), 1)
		self.assertEqual(Log.objects.filter(adherente=self.inactivo).count(), 0)

	def test_sin_permiso_no_procesa(self):
		self.client.force_login(self.creador)
		respuesta = self.subir([[self.titular.cbu]])
		self.assertEqual(respuesta.status_code, 403)
		self.titular.refresh_from_db()
		self.assertTrue(self.titular.is_active)

	def test_solapa_visible_solo_con_permiso(self):
		respuesta = self.client.get(reverse('stats'))
		self.assertContains(respuesta, 'Bajas masivas')
		self.assertContains(respuesta, reverse('bajas_masivas'))
		self.assertContains(respuesta, 'window.confirm')
		self.client.force_login(self.creador)
		self.creador.user_permissions.add(Permission.objects.get(codename='can_view_stats', content_type__model='titular'))
		respuesta = self.client.get(reverse('stats'))
		self.assertNotContains(respuesta, 'Bajas masivas')

	def test_archivo_invalido_no_modifica_registros(self):
		archivo = SimpleUploadedFile('cbus.xlsx', b'no es excel')
		respuesta = self.client.post(reverse('bajas_masivas'), {'archivo': archivo})
		self.assertEqual(respuesta.status_code, 400)
		self.titular.refresh_from_db()
		self.assertTrue(self.titular.is_active)

	def test_error_en_una_fila_revierte_y_continua(self):
		crear_log = Log.objects.create
		llamadas = 0

		def registrar(*args, **kwargs):
			nonlocal llamadas
			llamadas += 1
			if llamadas == 1:
				raise DatabaseError('fallo de prueba')
			return crear_log(*args, **kwargs)

		with patch.object(Log.objects, 'create', side_effect=registrar):
			respuesta = self.subir([[self.titular.cbu], [self.titular.cbu]])

		filas = list(load_workbook(BytesIO(respuesta.content)).active.values)
		self.assertEqual([fila[2] for fila in filas[1:]], ['Error', 'Baja realizada'])
		self.assertEqual(Log.objects.filter(adherente=self.activo, movimiento='Baja').count(), 1)
		self.titular.refresh_from_db()
		self.assertFalse(self.titular.is_active)
