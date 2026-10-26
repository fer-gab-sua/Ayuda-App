from datetime import datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.admin.sites import AdminSite
from django.contrib.auth.models import Permission, User
from django.db import DatabaseError
from django.test import RequestFactory, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .admin import AdherenteAdmin, MovimientoAdmin, TitularAdmin
from .forms import AdherenteForm, ClientForm
from .models import Adherente, DatosUser, Movimiento, Sucursales, Titular
from .movimientos import guardar_con_movimiento


class MovimientosTests(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user(username='vendedor', password='clave')
        self.otro = User.objects.create_user(username='supervisor', password='clave')
        self.sucursal = Sucursales.objects.create(descripcion='Centro')
        DatosUser.objects.create(user=self.usuario, legajo=12, sucursal=self.sucursal)
        DatosUser.objects.create(user=self.otro, legajo=13, sucursal=self.sucursal)
        self.usuario.user_permissions.add(Permission.objects.get(
            codename='can_view_stats', content_type__model='titular',
        ))
        self.titular = Titular.objects.create(
            cbu='0012345678901234567890', name='Ana', last_name='Perez',
            document_type='DNI', document='12345678', sex='F',
            street_address='Calle', number='1', province='Buenos Aires',
            city='Ciudad', postal_code='1000', user_upload=self.usuario, plan='Basico',
        )
        self.adherente = Adherente.objects.create(
            titular=self.titular, name='Maria', last_name='Perez',
            document_type='DNI', document='22345678', sex='F',
            street_address='Calle', number='1', province='Buenos Aires',
            city='Ciudad', postal_code='1000', user_upload=self.usuario,
            sucursal='Centro', legajo=12, plan='Basico',
        )
        self.client.force_login(self.usuario)
        self.fecha = datetime(2026, 10, 13, 12, 0, 0)

    def guardar(self, persona, usuario=None, fecha=None):
        with patch('pianoadherentes.movimientos.timezone.now', return_value=fecha or self.fecha):
            return guardar_con_movimiento(persona, usuario or self.usuario)

    def datos(self, persona, plan='Plus', **cambios):
        form = ClientForm(instance=persona) if isinstance(persona, Titular) else AdherenteForm(instance=persona)
        datos = {nombre: getattr(persona, nombre) or '' for nombre in form.fields}
        datos['plan'] = plan
        datos.update(cambios)
        return datos

    def filtros(self, **cambios):
        return {
            'init_date': '2026-10-13', 'end_date': '2026-10-13',
            'listado': 'todos', **cambios,
        }

    def test_altas_plus_titular_y_adherente(self):
        self.titular.pk = None
        self.titular._state.adding = True
        self.titular.cbu = '9' * 22
        self.titular.plan = 'Plus'
        alta = self.guardar(self.titular)
        self.assertEqual(alta.movimiento, 'Alta')
        self.assertEqual(alta.tipo_persona, 'Titular')
        self.assertEqual(alta.plan_anterior, '')
        self.assertTrue(alta.venta_plus)
        self.assertFalse(alta.reingreso_plus)
        self.adherente.pk = None
        self.adherente._state.adding = True
        self.adherente.plan = 'Plus'
        alta = self.guardar(self.adherente)
        self.assertEqual(alta.tipo_persona, 'Adherente')
        self.assertTrue(alta.venta_plus)
        self.assertEqual(alta.user, self.usuario)
        self.assertEqual(alta.legajo, 12)
        self.assertEqual(alta.sucursal, 'Centro')

    def test_cambio_de_plan_y_edicion_plus_no_duplica_venta(self):
        for persona in (self.titular, self.adherente):
            with self.subTest(persona=type(persona).__name__):
                persona.plan = 'Plus'
                cambio = self.guardar(persona, self.otro)
                self.assertEqual(cambio.plan_anterior, 'Basico')
                self.assertEqual(cambio.plan_nuevo, 'Plus')
                self.assertEqual(cambio.user, self.otro)
                self.assertTrue(cambio.venta_plus)
                persona.phone = '999'
                edicion = self.guardar(persona)
                self.assertFalse(edicion.venta_plus)
                self.assertEqual(edicion.plan_anterior, 'Plus')
        self.assertEqual(Movimiento.objects.filter(venta_plus=True).count(), 2)

    def test_baja_y_reactivacion_plus_con_fechas_y_usuario(self):
        for persona in (self.titular, self.adherente):
            with self.subTest(persona=type(persona).__name__):
                persona.plan = 'Plus'
                ingreso = self.guardar(persona, fecha=self.fecha - timedelta(days=10))
                persona.is_active = False
                baja = self.guardar(persona, self.otro, fecha=self.fecha - timedelta(seconds=2))
                self.assertEqual(baja.movimiento, 'Baja')
                self.assertTrue(baja.salida_plus)
                self.assertFalse(baja.venta_plus)
                persona.is_active = True
                vuelta = self.guardar(persona)
                self.assertEqual(vuelta.movimiento, 'Reactivacion')
                self.assertTrue(vuelta.venta_plus)
                self.assertTrue(vuelta.reingreso_plus)
                self.assertEqual(vuelta.ingreso_plus_anterior, ingreso.created)
                self.assertEqual(vuelta.salida_plus_anterior, baja.created)
                self.assertEqual(vuelta.usuario_salida_plus, self.otro.username)
                self.assertAlmostEqual(vuelta.dias_fuera_plus, 2 / 86400)
                self.assertIn('Tiempo fuera de Plus: 0.000023 dias.', vuelta.observaciones)
                self.assertIn('Revisar antes de liquidar', vuelta.observaciones)
                self.assertIn('supervisor', vuelta.observaciones)
                persona.refresh_from_db()
                self.assertIsNone(persona.deleted)
                repetido = self.guardar(persona)
                self.assertFalse(repetido.venta_plus)
                self.assertFalse(repetido.reingreso_plus)

    def test_salida_por_cambio_de_plan_y_reingreso(self):
        self.adherente.plan = 'Plus'
        ingreso = self.guardar(self.adherente, fecha=self.fecha - timedelta(days=4))
        self.adherente.plan = 'Basico'
        salida = self.guardar(self.adherente, fecha=self.fecha - timedelta(days=1))
        self.assertTrue(salida.salida_plus)
        self.assertEqual(salida.movimiento, 'Modificacion')
        self.adherente.plan = 'Plus'
        vuelta = self.guardar(self.adherente)
        self.assertTrue(vuelta.reingreso_plus)
        self.assertEqual(vuelta.ingreso_plus_anterior, ingreso.created)
        self.assertEqual(vuelta.salida_plus_anterior, salida.created)
        self.assertEqual(vuelta.dias_fuera_plus, 1)

    def test_persona_inactiva_no_comisiona_hasta_reactivacion(self):
        self.adherente.is_active = False
        self.guardar(self.adherente)
        self.adherente.plan = 'Plus'
        self.assertFalse(self.guardar(self.adherente).venta_plus)
        self.adherente.phone = '999'
        self.assertFalse(self.guardar(self.adherente).venta_plus)
        self.adherente.is_active = True
        activacion = self.guardar(self.adherente)
        self.assertTrue(activacion.venta_plus)
        self.assertFalse(activacion.reingreso_plus)

    def test_plus_preexistente_no_inventa_fecha_de_ingreso(self):
        self.adherente.plan = 'Plus'
        self.adherente.is_active = False
        self.adherente.deleted = self.fecha - timedelta(days=1)
        self.adherente.save()
        self.adherente.phone = '999'
        self.guardar(self.adherente)
        self.adherente.is_active = True
        vuelta = self.guardar(self.adherente)
        self.assertTrue(vuelta.reingreso_plus)
        self.assertIsNone(vuelta.ingreso_plus_anterior)
        self.assertEqual(vuelta.salida_plus_anterior, self.fecha - timedelta(days=1))
        self.assertEqual(vuelta.usuario_salida_plus, '')
        self.assertIn('anterior al historial', vuelta.observaciones)

    def test_perfil_faltante_queda_aclarado(self):
        sin_perfil = User.objects.create_user(username='sin-perfil')
        self.titular.plan = 'Plus'
        movimiento = self.guardar(self.titular, sin_perfil)
        self.assertEqual(movimiento.user, sin_perfil)
        self.assertIsNone(movimiento.legajo)
        self.assertEqual(movimiento.sucursal, '')
        self.assertIn('sin legajo/sucursal', movimiento.observaciones)

    def test_datos_historicos_no_cambian_con_persona_o_usuario(self):
        self.titular.plan = 'Plus'
        ingreso = self.guardar(self.titular)
        self.titular.name = 'Nuevo nombre'
        self.titular.plan = 'Basico'
        self.titular.save()
        self.usuario.username = 'Otro nombre'
        self.usuario.save()
        self.sucursal.descripcion = 'Nueva sucursal'
        self.sucursal.save()
        ingreso.refresh_from_db()
        self.assertEqual(ingreso.nombre, 'Ana Perez')
        self.assertEqual(ingreso.usuario, 'vendedor')
        self.assertEqual(ingreso.sucursal, 'Centro')
        self.assertEqual(ingreso.plan_nuevo, 'Plus')

    def test_error_del_historial_revierte_persona(self):
        self.titular.plan = 'Plus'
        with patch.object(Movimiento.objects, 'create', side_effect=DatabaseError('fallo')):
            with self.assertRaises(DatabaseError):
                self.guardar(self.titular)
        self.titular.refresh_from_db()
        self.assertEqual(self.titular.plan, 'Basico')
        self.assertEqual(Movimiento.objects.count(), 0)

    def test_altas_desde_vistas(self):
        datos = self.datos(self.titular, cbu='9' * 22)
        respuesta = self.client.post(reverse('create_client'), datos)
        self.assertEqual(respuesta.status_code, 200)
        alta = Movimiento.objects.get(tipo_persona='Titular')
        self.assertEqual(alta.movimiento, 'Alta')
        self.assertTrue(alta.venta_plus)
        respuesta = self.client.post(reverse('create_adherente'), {
            **self.datos(self.adherente), 'titular_id': self.titular.pk,
        })
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(Movimiento.objects.get(tipo_persona='Adherente').venta_plus)

    def test_edicion_desde_todas_las_vistas(self):
        for ruta in ('update_titular', 'client_detail'):
            with self.subTest(ruta=ruta):
                self.titular.plan = 'Basico'
                self.titular.save()
                respuesta = self.client.post(reverse(ruta, args=[self.titular.pk]), self.datos(self.titular))
                self.assertIn(respuesta.status_code, (200, 302))
                cambio = Movimiento.objects.latest('pk')
                self.assertEqual(cambio.plan_anterior, 'Basico')
                self.assertTrue(cambio.venta_plus)
        respuesta = self.client.post(reverse('update_adherente', args=[self.adherente.pk]), self.datos(self.adherente))
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(Movimiento.objects.latest('pk').venta_plus)

    def test_baja_y_reactivacion_adherente_desde_vistas(self):
        self.adherente.plan = 'Plus'
        self.guardar(self.adherente)
        respuesta = self.client.post(reverse('baja_adherente', args=[self.adherente.pk]))
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(Movimiento.objects.latest('pk').salida_plus)
        respuesta = self.client.post(reverse('reactiv_adherente', args=[self.adherente.pk]))
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(Movimiento.objects.latest('pk').reingreso_plus)
        self.client.post(reverse('reactiv_adherente', args=[self.adherente.pk]))
        self.assertFalse(Movimiento.objects.latest('pk').venta_plus)

    def test_baja_de_titular_registra_todo_el_grupo(self):
        self.titular.plan = 'Plus'
        self.adherente.plan = 'Plus'
        self.guardar(self.titular)
        self.guardar(self.adherente)
        respuesta = self.client.post(reverse('client_baja', args=[self.titular.pk]))
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(Movimiento.objects.filter(movimiento='Baja', salida_plus=True).count(), 2)

    def test_formulario_invalido_no_registra(self):
        respuesta = self.client.post(
            reverse('update_titular', args=[self.titular.pk]),
            self.datos(self.titular, document='invalido'),
        )
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, 'DNI')
        self.assertEqual(Movimiento.objects.count(), 0)
        self.titular.refresh_from_db()
        self.assertEqual(self.titular.plan, 'Basico')

    def test_admin_registra_y_no_permite_editar_historial(self):
        peticion = RequestFactory().post('/admin/')
        peticion.user = self.usuario
        for persona, clase in ((self.titular, TitularAdmin), (self.adherente, AdherenteAdmin)):
            persona.plan = 'Plus'
            clase(type(persona), AdminSite()).save_model(peticion, persona, None, True)
            self.assertTrue(Movimiento.objects.latest('pk').venta_plus)
        admin = MovimientoAdmin(Movimiento, AdminSite())
        self.assertFalse(admin.has_add_permission(peticion))
        self.assertFalse(admin.has_change_permission(peticion))
        self.assertFalse(admin.has_delete_permission(peticion))

    def test_reporte_incluye_ultimo_dia_y_excluye_siguiente(self):
        self.titular.plan = 'Plus'
        self.guardar(self.titular, fecha=datetime(2026, 10, 13, 23, 59, 59))
        self.adherente.plan = 'Plus'
        self.guardar(self.adherente, fecha=datetime(2026, 10, 14, 0, 0, 0))
        respuesta = self.client.get(reverse('movimientos'), self.filtros())
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.context['pagina'].paginator.count, 1)
        self.assertEqual(respuesta.context['ventas'], 1)
        self.assertContains(respuesta, '23:59:59')

    def test_reporte_y_excel_reingreso_con_historia_fuera_del_periodo(self):
        self.adherente.plan = 'Plus'
        ingreso = self.guardar(self.adherente, fecha=self.fecha - timedelta(days=5))
        self.adherente.is_active = False
        baja = self.guardar(self.adherente, self.otro, fecha=self.fecha - timedelta(days=1))
        self.adherente.is_active = True
        self.guardar(self.adherente)
        filtros = self.filtros(listado='reingresos')
        respuesta = self.client.get(reverse('movimientos'), filtros)
        self.assertEqual(respuesta.context['ventas'], 1)
        self.assertEqual(respuesta.context['reingresos'], 1)
        self.assertContains(respuesta, '08/10/2026')
        self.assertContains(respuesta, '12/10/2026')
        respuesta = self.client.get(reverse('movimientos'), {**filtros, 'exportar': '1'})
        libro = load_workbook(BytesIO(respuesta.content))
        filas = list(libro['Movimientos'].values)
        datos = dict(zip(filas[0], filas[1]))
        self.assertEqual(datos['Ingreso Plus anterior'], ingreso.created.strftime('%d/%m/%Y %H:%M:%S'))
        self.assertEqual(datos['Baja o salida Plus anterior'], baja.created.strftime('%d/%m/%Y %H:%M:%S'))
        self.assertEqual(datos['Dias fuera de Plus'], 1)
        self.assertNotIn('Segundos fuera de Plus', datos)
        self.assertEqual(datos['Usuario baja o salida anterior'], 'supervisor')
        self.assertEqual(list(libro['Resumen por vendedor'].values)[1][4:], (1, 1))

    def test_dias_fuera_plus_fraccionarios_en_pantalla_y_excel(self):
        self.adherente.plan = 'Plus'
        self.guardar(self.adherente, fecha=self.fecha - timedelta(days=2))
        self.adherente.is_active = False
        self.guardar(self.adherente, fecha=self.fecha - timedelta(hours=12))
        self.adherente.is_active = True
        vuelta = self.guardar(self.adherente)
        self.assertEqual(vuelta.dias_fuera_plus, 0.5)
        self.assertIn('Tiempo fuera de Plus: 0.500000 dias.', vuelta.observaciones)
        respuesta = self.client.get(reverse('movimientos'), self.filtros(listado='reingresos'))
        self.assertContains(respuesta, 'Dias fuera de Plus')
        self.assertContains(respuesta, '<td>0.5</td>', html=True)
        self.assertNotContains(respuesta, 'Segundos fuera de Plus')
        respuesta = self.client.get(reverse('movimientos'), self.filtros(listado='reingresos', exportar='1'))
        hoja = load_workbook(BytesIO(respuesta.content))['Movimientos']
        self.assertEqual(hoja['W2'].value, 0.5)
        self.assertEqual(hoja['W2'].data_type, 'n')
        self.assertEqual(hoja['W2'].number_format, '0.######')

    def test_dias_fuera_plus_cero_y_fecha_desconocida(self):
        self.adherente.plan = 'Plus'
        ingreso = self.guardar(self.adherente)
        self.assertIsNone(ingreso.dias_fuera_plus)
        self.adherente.is_active = False
        self.guardar(self.adherente)
        self.adherente.is_active = True
        vuelta = self.guardar(self.adherente)
        self.assertEqual(vuelta.dias_fuera_plus, 0)
        respuesta = self.client.get(reverse('movimientos'), self.filtros(listado='reingresos'))
        self.assertContains(respuesta, '<td>0</td>', html=True)

    def test_estadisticas_muestran_movimientos_de_todos_los_usuarios(self):
        self.adherente.plan = 'Plus'
        propio = self.guardar(self.adherente)
        self.titular.plan = 'Plus'
        ajeno = self.guardar(self.titular, self.otro)
        respuesta = self.client.get(reverse('movimientos'), self.filtros())
        self.assertEqual(respuesta.context['pagina'].paginator.count, 2)
        self.assertNotContains(respuesta, 'Solo se muestran los movimientos realizados por tu usuario.')
        excel = self.client.get(reverse('movimientos'), self.filtros(exportar='1'))
        hoja = load_workbook(BytesIO(excel.content))['Movimientos']
        self.assertEqual(hoja.max_row, 3)
        self.assertEqual({hoja['A2'].value, hoja['A3'].value}, {propio.pk, ajeno.pk})
        respuesta = self.client.get(reverse('movimientos'), self.filtros(usuario=self.otro.pk))
        self.assertEqual(respuesta.context['pagina'].paginator.count, 1)
        self.assertEqual([registro.pk for registro in respuesta.context['pagina']], [ajeno.pk])
        self.client.force_login(self.otro)
        self.assertEqual(self.client.get(reverse('movimientos')).status_code, 403)

    def test_filtros_de_persona_operacion_usuario_y_sucursal(self):
        self.titular.plan = 'Plus'
        self.guardar(self.titular)
        self.adherente.plan = 'Plus'
        self.guardar(self.adherente, self.otro)
        respuesta = self.client.get(reverse('movimientos'), self.filtros(
            usuario=self.usuario.pk, sucursal='Centro', tipo_persona='Titular',
            movimiento='Modificacion', listado='plus',
        ))
        self.assertEqual(respuesta.context['pagina'].paginator.count, 1)
        self.assertEqual(respuesta.context['ventas'], 1)

    def test_reporte_rechaza_fechas_y_filtros_invalidos(self):
        for filtros in (
            self.filtros(init_date='invalido'),
            self.filtros(end_date='2026-10-12'),
            self.filtros(tipo_persona='invalido'),
            self.filtros(listado='invalido'),
        ):
            with self.subTest(filtros=filtros):
                respuesta = self.client.get(reverse('movimientos'), filtros)
                self.assertEqual(respuesta.status_code, 400)

    def test_excel_exporta_todas_las_paginas_y_conserva_textos(self):
        self.titular.name = '=1+1'
        self.titular.plan = 'Plus'
        for _ in range(51):
            self.guardar(self.titular)
        respuesta = self.client.get(reverse('movimientos'), self.filtros())
        self.assertEqual(len(respuesta.context['pagina']), 50)
        respuesta = self.client.get(reverse('movimientos'), self.filtros(exportar='1'))
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro['Movimientos'].max_row, 52)
        self.assertEqual(libro['Movimientos']['F2'].data_type, 's')
        self.assertEqual(libro['Movimientos']['H2'].value, self.titular.cbu)
        self.assertEqual(libro['Movimientos']['H2'].data_type, 's')
        self.assertEqual(list(libro['Resumen por vendedor'].values)[1][4], 1)
