from django.contrib.auth.decorators import permission_required
from django.contrib.auth.models import User
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from openpyxl import Workbook

from ..forms import MovimientosForm
from ..models import Movimiento


@permission_required('pianoadherentes.can_view_stats', raise_exception=True)
def movimientos(request):
    registros = Movimiento.objects.all()
    sucursales = registros.order_by('sucursal').values_list('sucursal', flat=True).distinct()
    fecha = timezone.now().date()
    form = MovimientosForm(
        request.GET if request.GET else None,
        initial={'init_date': fecha.replace(day=1), 'end_date': fecha, 'listado': 'todos'},
        usuarios=User.objects.order_by('username'), sucursales=sucursales,
    )
    contexto = {'form': form}
    if not form.is_bound:
        return render(request, 'movimientos.html', contexto)
    if not form.is_valid():
        return render(request, 'movimientos.html', contexto, status=400)

    datos = form.cleaned_data
    registros = registros.filter(
        created__date__gte=datos['init_date'], created__date__lte=datos['end_date'],
    )
    for campo in ('usuario', 'sucursal', 'tipo_persona', 'movimiento'):
        valor = datos.get(campo)
        if valor:
            registros = registros.filter(**{'user' if campo == 'usuario' else campo: valor})
    if datos['listado'] == 'plus':
        registros = registros.filter(venta_plus=True)
    elif datos['listado'] == 'reingresos':
        registros = registros.filter(reingreso_plus=True)
    resumen = list(registros.order_by().values('user_id', 'usuario', 'legajo', 'sucursal').annotate(
        ventas=Count('pk', filter=Q(venta_plus=True)),
        reingresos=Count('pk', filter=Q(reingreso_plus=True)),
    ).filter(ventas__gt=0).order_by('usuario', 'sucursal', 'legajo'))
    if request.GET.get('exportar') == '1':
        return exportar_movimientos(registros, resumen)

    parametros = request.GET.copy()
    parametros.pop('page', None)
    parametros.pop('exportar', None)
    contexto.update({
        'consultado': True,
        'pagina': Paginator(registros.order_by('-created', '-pk'), 50).get_page(request.GET.get('page')),
        'resumen': resumen, 'consulta': parametros.urlencode(),
        'ventas': sum(fila['ventas'] for fila in resumen),
        'reingresos': sum(fila['reingresos'] for fila in resumen),
    })
    return render(request, 'movimientos.html', contexto)


def exportar_movimientos(registros, resumen):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Movimientos'
    hoja.append([
        'Movimiento ID', 'Fecha y hora', 'Tipo persona', 'Titular ID', 'Adherente ID',
        'Nombre', 'Documento', 'CBU', 'Operacion', 'Plan anterior', 'Plan nuevo',
        'Activo anterior', 'Activo nuevo', 'Usuario ID', 'Usuario', 'Legajo', 'Sucursal',
        'Venta Plus', 'Reingreso Plus', 'Ingreso Plus anterior', 'Baja o salida Plus anterior',
        'Usuario baja o salida anterior', 'Dias fuera de Plus', 'Observaciones',
    ])
    for registro in registros.iterator():
        hoja.append([
            registro.pk, registro.created.strftime('%d/%m/%Y %H:%M:%S'),
            registro.tipo_persona, registro.titular_id, registro.adherente_id,
            registro.nombre, registro.documento, registro.cbu, registro.movimiento,
            registro.plan_anterior, registro.plan_nuevo,
            '' if registro.activo_anterior is None else ('Si' if registro.activo_anterior else 'No'),
            'Si' if registro.activo_nuevo else 'No', registro.user_id, registro.usuario,
            registro.legajo, registro.sucursal, 'Si' if registro.venta_plus else 'No',
            'Si' if registro.reingreso_plus else 'No',
            registro.ingreso_plus_anterior.strftime('%d/%m/%Y %H:%M:%S') if registro.ingreso_plus_anterior else '',
            registro.salida_plus_anterior.strftime('%d/%m/%Y %H:%M:%S') if registro.salida_plus_anterior else '',
            registro.usuario_salida_plus, registro.dias_fuera_plus, registro.observaciones,
        ])
        hoja.cell(row=hoja.max_row, column=23).number_format = '0.######'
    totales = libro.create_sheet('Resumen por vendedor')
    totales.append(['Usuario ID', 'Usuario', 'Legajo', 'Sucursal', 'Ventas Plus', 'Reingresos incluidos (revisar)'])
    for fila in resumen:
        totales.append([fila['user_id'], fila['usuario'], fila['legajo'], fila['sucursal'], fila['ventas'], fila['reingresos']])
    totales.append(['', 'TOTAL', '', '', sum(f['ventas'] for f in resumen), sum(f['reingresos'] for f in resumen)])
    for tabla in (hoja, totales):
        tabla.freeze_panes = 'A2'
        tabla.auto_filter.ref = tabla.dimensions
        for fila in tabla:
            for celda in fila:
                if isinstance(celda.value, str):
                    celda.data_type = 's'
    respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = 'attachment; filename="movimientos_plus.xlsx"'
    libro.save(respuesta)
    return respuesta
